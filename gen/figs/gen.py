import os
OUTDIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'out')   # 中間生成物（xlsx）の置き場。git 管理外
os.makedirs(OUTDIR, exist_ok=True)
import openpyxl
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.properties import PageSetupProperties

STEP = 10
COLORS = ["FFC000", "FFFF00", "92D050", "00B050"]
GRAY = "D9D9D9"


def simulate(cap, make_fn, use_ms, tmax):
    """C++ と同じ意味論：作る側 = 作り終えてから push（満杯なら待つ）、使う側 = pop してから使う"""
    q = []
    i = 0
    ready = make_fn(0)
    p_start = 0
    blocked_from = None
    busy_end, cur = None, None
    idle_from = 0
    push, pop, qv = {}, {}, {}
    prod, cons = [], []

    for t in range(0, tmax + 1):
        if busy_end == t:
            cons.append((busy_end - use_ms, t, "#%d" % cur, "item"))
            busy_end, cur = None, None
            idle_from = t

        if blocked_from is None and t == ready:
            prod.append((p_start, t, "#%d" % i, "item"))
            if len(q) < cap:
                q.append(i); push[t] = i; i += 1
                p_start = t; ready = t + make_fn(i)
            else:
                blocked_from = t

        if busy_end is None and q:
            v = q.pop(0); pop[t] = v
            if t > idle_from:
                cons.append((idle_from, t, "idle", "wait"))
            busy_end = t + use_ms; cur = v

        if blocked_from is not None and len(q) < cap:
            prod.append((blocked_from, t, "wait", "wait"))
            q.append(i); push[t] = i; i += 1
            p_start = t; ready = t + make_fn(i); blocked_from = None

        qv[t] = len(q)

    if busy_end is not None:
        cons.append((busy_end - use_ms, tmax, "#%d" % cur, "item"))
    if blocked_from is not None:
        prod.append((blocked_from, tmax, "wait", "wait"))
    else:
        prod.append((p_start, tmax, "#%d" % i, "item"))
    return push, pop, qv, prod, cons


def put_band(ws, row, spans, times):
    ci = {t: 2 + k for k, t in enumerate(times)}
    for (a, b, label, kind) in spans:
        lo, hi = a + STEP, b
        lo = max(lo, times[0]); hi = min(hi, times[-1])
        if lo > hi:
            continue
        color = GRAY if kind == "wait" else COLORS[int(label[1:]) % len(COLORS)]
        c1, c2 = ci[lo], ci[hi]
        if c2 > c1:
            ws.merge_cells(start_row=row, start_column=c1, end_row=row, end_column=c2)
        cell = ws.cell(row=row, column=c1, value=label)
        cell.alignment = Alignment(horizontal="center")
        for c in range(c1, c2 + 1):
            ws.cell(row=row, column=c).fill = PatternFill("solid", fgColor=color)


def block(ws, row0, title, cap, make_fn, use_ms, times):
    ws.cell(row=row0, column=1, value=title).font = Font(bold=True, size=12)
    push, pop, qv, prod, cons = simulate(cap, make_fn, use_ms, times[-1] + 300)
    for k, lbl in enumerate(["Time [mS]", "push #", "pop #", "Q value", "producer", "consumer"]):
        ws.cell(row=row0 + 1 + k, column=1, value=lbl)
    last = None
    for k, t in enumerate(times):
        c = 2 + k
        ws.cell(row=row0 + 1, column=c, value=t).alignment = Alignment(horizontal="center")
        if t in push:
            last = push[t]
        if last is not None:
            ws.cell(row=row0 + 2, column=c, value=last).alignment = Alignment(horizontal="center")
        if t in pop:
            ws.cell(row=row0 + 3, column=c, value=pop[t]).alignment = Alignment(horizontal="center")
        ws.cell(row=row0 + 4, column=c, value=qv[t]).alignment = Alignment(horizontal="center")
    for c in range(1, 2 + len(times)):
        ws.cell(row=row0 + 1, column=c).border = Border(bottom=Side(style="medium"))
    put_band(ws, row0 + 5, prod, times)
    put_band(ws, row0 + 6, cons, times)


def sheet(path, cap_a, cap_b, make_fn, use_ms, times):
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "timing"
    block(ws, 1, "Capacity %d" % cap_a, cap_a, make_fn, use_ms, times)
    block(ws, 10, "Capacity %d" % cap_b, cap_b, make_fn, use_ms, times)
    ws.column_dimensions['A'].width = 15
    for k in range(len(times)):
        ws.column_dimensions[get_column_letter(2 + k)].width = 6.2
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.page_margins.left = ws.page_margins.right = 0.2
    ws.page_margins.top = ws.page_margins.bottom = 0.2
    wb.save(path)


TIMES = list(range(10, 200, 10))
sheet(os.path.join(OUTDIR, 'qA2.xlsx'), 1, 4, lambda i: 10, 50, TIMES)
sheet(os.path.join(OUTDIR, 'qB2.xlsx'), 1, 4, lambda i: 90 if i % 5 == 4 else 10, 30, TIMES)
print('saved')
