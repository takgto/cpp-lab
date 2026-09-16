import os
OUTDIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'out')   # 中間生成物（xlsx）の置き場。git 管理外
os.makedirs(OUTDIR, exist_ok=True)
# 3段パイプラインのタイムチャートを ex03 と同じ表現で描く
import openpyxl
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.properties import PageSetupProperties

STEP = 10
COLORS = ["FFC000", "FFFF00", "92D050", "00B050", "00B0F0", "E8A2A2",
          "B4A7D6", "8EA9DB"]
GRAY = "D9D9D9"      # 上流が空で手待ち（idle）
GRAY2 = "BFBFBF"     # 下流が満杯で入れられない（wait）


# ------------------------------------------------------------------
# 1ms 刻みの素朴なシミュレータ。C++ と同じ意味論：
#   pop（空なら待つ） → 処理する → push（満杯なら待つ）
# ------------------------------------------------------------------
class Q:
    def __init__(self, cap): self.cap = cap; self.v = []


class W:
    """1本のスレッド。q_in から取り出し、dur(f) ms 処理し、q_out に入れる"""
    def __init__(self, name, q_in, q_out, dur, items):
        self.name, self.qi, self.qo, self.dur = name, q_in, q_out, dur
        self.items = list(items)          # この人が担当するフレーム番号（Read は生成）
        self.st = 'pop'                   # pop / work / push
        self.cur = None
        self.end = 0
        self.mark = 0                     # 現在の区間の開始時刻
        self.spans = []

    def close(self, t, label, kind):
        if t > self.mark:
            self.spans.append((self.mark, t, label, kind))
        self.mark = t

    def step(self, t, ev_push, ev_pop):
        if self.st == 'work' and t >= self.end:
            self.close(t, "#%d" % self.cur, 'item')
            self.st = 'push' if self.qo is not None else 'pop'
            if self.qo is None:
                self.cur = None
        if self.st == 'push':
            if len(self.qo.v) < self.qo.cap:
                self.close(t, 'wait', 'wait2')
                self.qo.v.append(self.cur); ev_push.setdefault(t, self.cur)
                self.cur = None; self.st = 'pop'
            else:
                return
        if self.st == 'pop':
            if self.qi is None:                       # Read：待たずに次を作る
                if not self.items: return
                self.close(t, 'idle', 'wait')
                self.cur = self.items.pop(0)
                self.st = 'work'; self.end = t + self.dur(self.cur)
            elif self.qi.v:
                self.close(t, 'idle', 'wait')
                self.cur = self.qi.v.pop(0); ev_pop.setdefault(t, self.cur)
                self.st = 'work'; self.end = t + self.dur(self.cur)


def run(workers, queues, tmax):
    qlog = {q: {} for q in queues}
    pushes, pops = {}, {}
    for t in range(0, tmax + 1):
        for w in reversed(workers):          # 下流から動かす（同じ時刻の pop を先に）
            w.step(t, pushes, pops)
        for q in queues:
            qlog[q][t] = len(q.v)
    for w in workers:
        if w.st == 'work':   w.close(tmax, "#%d" % w.cur, 'item')
        elif w.st == 'push': w.close(tmax, 'wait', 'wait2')
        else:                w.close(tmax, 'idle', 'wait')
    return qlog, pushes, pops


# ------------------------------------------------------------------ 描画
def put_band(ws, row, spans, times):
    ci = {t: 2 + k for k, t in enumerate(times)}
    for (a, b, label, kind) in spans:
        lo = (a // STEP + 1) * STEP          # a より後の最初の格子点
        hi = (b // STEP) * STEP              # b 以下の最後の格子点
        lo = max(lo, times[0]); hi = min(hi, times[-1])
        if lo > hi:
            continue
        color = GRAY if kind == 'wait' else (GRAY2 if kind == 'wait2'
                                             else COLORS[int(label[1:]) % len(COLORS)])
        c1, c2 = ci[lo], ci[hi]
        if c2 > c1:
            ws.merge_cells(start_row=row, start_column=c1, end_row=row, end_column=c2)
        cell = ws.cell(row=row, column=c1, value=label)
        cell.alignment = Alignment(horizontal='center')
        for c in range(c1, c2 + 1):
            ws.cell(row=row, column=c).fill = PatternFill('solid', fgColor=color)


def block(ws, row0, title, workers, queues, qnames, times, note=None):
    ws.cell(row=row0, column=1, value=title).font = Font(bold=True, size=12)
    if note:
        ws.cell(row=row0, column=16, value=note).font = Font(size=10)
    qlog, _, _ = run(workers, queues, times[-1] + STEP)

    r = row0 + 1
    ws.cell(row=r, column=1, value='Time [mS]')
    for k, t in enumerate(times):
        ws.cell(row=r, column=2 + k, value=t).alignment = Alignment(horizontal='center')
    for c in range(1, 2 + len(times)):
        ws.cell(row=r, column=c).border = Border(bottom=Side(style='medium'))
    r += 1
    for q, nm in zip(queues, qnames):
        ws.cell(row=r, column=1, value=nm)
        for k, t in enumerate(times):
            ws.cell(row=r, column=2 + k, value=qlog[q][t]).alignment = Alignment(horizontal='center')
        r += 1
    for w in workers:
        ws.cell(row=r, column=1, value=w.name)
        put_band(ws, r, w.spans, times)
        r += 1
    return r


def finish(wb, ws, ncol, path):
    ws.column_dimensions['A'].width = 16
    for k in range(ncol):
        ws.column_dimensions[get_column_letter(2 + k)].width = 5.4
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.page_margins.left = ws.page_margins.right = 0.2
    ws.page_margins.top = ws.page_margins.bottom = 0.2
    wb.save(path)


READ, INFER, SHOW = 10, 60, 30
CAP = 4
TIMES = list(range(10, 241, 10))      # 24 列（文字を大きく見せるため範囲を短く）


def mk3(ninfer=1, nshow=1, read=READ, infer=INFER, show=SHOW, n=40, cap=CAP):
    q1, q2 = Q(cap), Q(cap)
    ws_ = [W('Read', None, q1, lambda f: read, range(n))]
    for k in range(ninfer):
        ws_.append(W('Infer' + (str(k + 1) if ninfer > 1 else ''), q1, q2, lambda f: infer, []))
    for k in range(nshow):
        ws_.append(W('Show' + (str(k + 1) if nshow > 1 else ''), q2, None, lambda f: show, []))
    return ws_, [q1, q2]


# ===================== 問3-1：直列 と 3段 =====================
wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'q3-1'

# 直列：1本のスレッドが read→infer→show を順にやる。3行に分けて描く
class Serial:
    def __init__(self, n):
        self.rows = {'Read': [], 'Infer': [], 'Show': []}
        t = 0
        for f in range(n):
            for nm, d in (('Read', READ), ('Infer', INFER), ('Show', SHOW)):
                self.rows[nm].append((t, t + d, '#%d' % f, 'item')); t += d


s = Serial(6)
ws.cell(row=1, column=1, value='構成① 直列 ―― 1本のスレッドが順にやる').font = Font(bold=True, size=12)
ws.cell(row=1, column=16, value='※ 3行に分けて描いているが、動いているのは常に1つだけ').font = Font(size=10)
ws.cell(row=2, column=1, value='Time [mS]')
for k, t in enumerate(TIMES):
    ws.cell(row=2, column=2 + k, value=t).alignment = Alignment(horizontal='center')
for c in range(1, 2 + len(TIMES)):
    ws.cell(row=2, column=c).border = Border(bottom=Side(style='medium'))
for i, nm in enumerate(['Read', 'Infer', 'Show']):
    ws.cell(row=3 + i, column=1, value=nm)
    put_band(ws, 3 + i, s.rows[nm], TIMES)

w, q = mk3()
block(ws, 8, '構成② 3段 ―― Read / Infer / Show を1本ずつ', w, q, ['q1 (R→I)', 'q2 (I→S)'], TIMES)
finish(wb, ws, len(TIMES), os.path.join(OUTDIR, 'p31.xlsx'))

# ===================== 問3-2：人を増やす =====================
wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'q3-2'
w, q = mk3(1, 1)
r = block(ws, 1, '構成② 3段（Infer 1人 / Show 1人）   一番遅い段 = Infer 60ms', w, q,
          ['q1 (R→I)', 'q2 (I→S)'], TIMES,
          note='※ コアが足りている場合の図（＝上限A だけが効く世界）')
w, q = mk3(2, 1)
r = block(ws, r + 2, '構成③ Infer 2人   一番遅い段 = Show 30ms', w, q,
          ['q1 (R→I)', 'q2 (I→S)'], TIMES)
finish(wb, ws, len(TIMES), os.path.join(OUTDIR, 'p32.xlsx'))

# ===================== 問3-3：キューの長さ =====================
wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'q3-3'
w, q = mk3(1, 1, 10, 60, 20, cap=8)
r = block(ws, 1, 'Infer が遅い（10 / 60 / 20 ms）', w, q, ['q1 (R→I)', 'q2 (I→S)'], TIMES)
w, q = mk3(1, 1, 10, 20, 60, cap=8)
r = block(ws, r + 2, 'Show が遅い（10 / 20 / 60 ms）', w, q, ['q1 (R→I)', 'q2 (I→S)'], TIMES)
w, q = mk3(1, 1, 60, 20, 10, cap=8)
block(ws, r + 2, 'Read が遅い（60 / 20 / 10 ms）', w, q, ['q1 (R→I)', 'q2 (I→S)'], TIMES)
finish(wb, ws, len(TIMES), os.path.join(OUTDIR, 'p33.xlsx'))

print('saved p31 / p32 / p33')
